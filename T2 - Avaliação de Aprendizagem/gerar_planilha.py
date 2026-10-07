# Gera indices_desempenho.xlsx: a planilha de cálculo dos índices do T2, com os
# dois modelos (atual e melhorado) simulados com a semente 1 e 100.000
# aleatórios, como no exemplo do M9. Cada fila tem uma aba no formato das
# tabelas do módulo (tempo por estado -> probabilidade, N, D, U e W), e todos os
# índices são fórmulas do Excel: mudando o tempo acumulado de um estado a
# planilha recalcula.
#
#   python3 gerar_planilha.py [--atual modelo_atual.yml]
#                             [--melhorado modelo_melhorado.yml]
#                             [--saida indices_desempenho.xlsx]
#
# Requisitos: pyyaml e openpyxl (pip install pyyaml openpyxl).
#
# Cores: azul = valor que veio do simulador ou do .yml (entrada), preto =
# fórmula, verde = ligação com outra aba, amarelo = parâmetro alterado na
# melhoria.

import argparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from analise_indices import analisar, kendall
from simulador_rede import carregar_modelo, simular_rede

SEEDS = [1, 2, 3, 4, 5]
NOMES = {"Q1": "N1 - Triagem", "Q2": "N2 - Especialista", "Q3": "N3 - Desenvolvimento"}

ENTRADA = Font(name="Arial", size=10, color="0000FF")
FORMULA = Font(name="Arial", size=10)
LIGACAO = Font(name="Arial", size=10, color="008000")
NEGRITO = Font(name="Arial", size=10, bold=True)
TITULO = Font(name="Arial", size=12, bold=True)
CABECALHO = PatternFill("solid", fgColor="D9E1F2")
ALTERADO = PatternFill("solid", fgColor="FFFF00")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
DIREITA = Alignment(horizontal="right")


def escreve(ws, ref, valor, fonte=FORMULA, formato=None, preenchimento=None):
    c = ws[ref]
    c.value = valor
    c.font = fonte
    if formato:
        c.number_format = formato
    if preenchimento:
        c.fill = preenchimento
    return c


def cabecalho(ws, linha, titulos, coluna_inicial=1):
    for k, t in enumerate(titulos):
        c = ws.cell(row=linha, column=coluna_inicial + k, value=t)
        c.font = NEGRITO
        c.fill = CABECALHO
        c.alignment = CENTRO


def larguras(ws, valores):
    for k, v in enumerate(valores, start=1):
        ws.column_dimensions[get_column_letter(k)].width = v


def aba_parametros(wb, modelos):
    # modelos: {nome: lista de Fila}. Devolve as referências de C, K e MU.
    ws = wb.create_sheet("Parâmetros")
    escreve(ws, "A1", "Parâmetros dos modelos (unidade de tempo: minutos)", TITULO)
    escreve(ws, "A2", "Azul = lido do .yml. MU é a taxa de atendimento de um servidor, "
                      "em clientes por hora: 60 / tempo médio de atendimento.")
    cabecalho(ws, 4, ["Modelo", "Fila", "Papel", "Kendall", "Servidores (C)",
                      "Capacidade (K)", "Atend. mín", "Atend. máx", "Tempo médio",
                      "MU (cl/h)", "Chegada mín", "Chegada máx", "1ª chegada"])
    refs = {}
    linha = 5
    for nome, filas in modelos.items():
        for f in filas:
            escreve(ws, f"A{linha}", nome)
            escreve(ws, f"B{linha}", f.id)
            escreve(ws, f"C{linha}", NOMES.get(f.id, f.id))
            escreve(ws, f"D{linha}", kendall(f.num_servidores, f.capacidade))
            escreve(ws, f"E{linha}", f.num_servidores, ENTRADA)
            escreve(ws, f"F{linha}", f.capacidade if f.capacidade is not None else "∞",
                    ENTRADA).alignment = DIREITA
            escreve(ws, f"G{linha}", f.atendimento_min, ENTRADA, "0.0")
            escreve(ws, f"H{linha}", f.atendimento_max, ENTRADA, "0.0")
            escreve(ws, f"I{linha}", f"=(G{linha}+H{linha})/2", formato="0.00")
            escreve(ws, f"J{linha}", f"=60/I{linha}", formato="0.0000")
            if f.chegada_min is not None:
                escreve(ws, f"K{linha}", f.chegada_min, ENTRADA, "0.0")
                escreve(ws, f"L{linha}", f.chegada_max, ENTRADA, "0.0")
                escreve(ws, f"M{linha}", f.primeira_chegada, ENTRADA, "0.0")
            refs[(nome, f.id)] = {"C": f"'Parâmetros'!$E${linha}",
                                  "K": f"'Parâmetros'!$F${linha}",
                                  "MU": f"'Parâmetros'!$J${linha}"}
            linha += 1
        linha += 1

    # roteamento: uma linha por aresta, mais a linha de saída (o que falta pra 1)
    linha += 1
    escreve(ws, f"A{linha}", "Roteamento entre as filas", NEGRITO)
    linha += 1
    nomes = list(modelos)
    cabecalho(ws, linha, ["Origem", "Destino"] + [f"Prob. {n.lower()}" for n in nomes])
    linha += 1
    base = {f.id: f for f in modelos[nomes[0]]}
    for origem in base:
        primeira = linha
        for destino in base[origem].destinos:
            escreve(ws, f"A{linha}", origem)
            escreve(ws, f"B{linha}", destino)
            valores = [{f.id: f for f in modelos[n]}[origem].destinos[destino] for n in nomes]
            for k, v in enumerate(valores):
                col = get_column_letter(3 + k)
                alterado = any(abs(v - w) > 1e-9 for w in valores)
                escreve(ws, f"{col}{linha}", v, ENTRADA, "0.00",
                        ALTERADO if alterado else None)
            linha += 1
        escreve(ws, f"A{linha}", origem)
        escreve(ws, f"B{linha}", "sai da rede (resolvido)")
        for k in range(len(nomes)):
            col = get_column_letter(3 + k)
            escreve(ws, f"{col}{linha}", f"=1-SUM({col}{primeira}:{col}{linha - 1})",
                    formato="0.00")
        linha += 1
    linha += 1
    escreve(ws, f"A{linha}", "Amarelo = parâmetro alterado na melhoria.")
    larguras(ws, [12, 24, 24, 12, 14, 14, 11, 11, 12, 11, 12, 12, 11])
    return refs


def aba_fila(wb, modelo, f, tempo_global, resultado, ref):
    # tabela por estado no formato do M9. Devolve os endereços dos índices.
    ws = wb.create_sheet(f"{modelo}_{f.id}")
    escreve(ws, "A1", f"{modelo} - {f.id} ({NOMES.get(f.id, f.id)})", TITULO)
    escreve(ws, "A3", "Kendall")
    escreve(ws, "B3", kendall(f.num_servidores, f.capacidade))
    escreve(ws, "A4", "Servidores (C)")
    escreve(ws, "B4", f"={ref['C']}", LIGACAO)
    escreve(ws, "A5", "Capacidade (K)")
    escreve(ws, "B5", f"={ref['K']}", LIGACAO).alignment = DIREITA
    escreve(ws, "A6", "MU (clientes/h)")
    escreve(ws, "B6", f"={ref['MU']}", LIGACAO, "0.0000")
    escreve(ws, "A7", "Tempo global (min)")
    escreve(ws, "B7", tempo_global, ENTRADA, "#,##0.0000")
    escreve(ws, "A8", "Perdas")
    escreve(ws, "B8", resultado["perdas"], ENTRADA, "#,##0")

    cabecalho(ws, 10, ["i", "Tempo acumulado (min)", "Probabilidade", "N (clientes)",
                       "D (cl/h)", "U"])
    primeira = 11
    for i, t in enumerate(resultado["times"]):
        r = primeira + i
        escreve(ws, f"A{r}", i).alignment = Alignment(horizontal="center")
        escreve(ws, f"B{r}", t, ENTRADA, "#,##0.0000")
        escreve(ws, f"C{r}", f"=B{r}/$B$7", formato="0.00%")
        escreve(ws, f"D{r}", f"=C{r}*A{r}", formato="0.0000")
        escreve(ws, f"E{r}", f"=C{r}*MIN(A{r},$B$4)*$B$6", formato="0.0000")
        escreve(ws, f"F{r}", f"=C{r}*MIN(A{r},$B$4)/$B$4", formato="0.0000")
    ultima = primeira + len(resultado["times"]) - 1
    tot = ultima + 1
    escreve(ws, f"A{tot}", "Total", NEGRITO)
    escreve(ws, f"B{tot}", f"=SUM(B{primeira}:B{ultima})", NEGRITO, "#,##0.0000")
    escreve(ws, f"C{tot}", f"=SUM(C{primeira}:C{ultima})", NEGRITO, "0.00%")
    escreve(ws, f"D{tot}", f"=SUM(D{primeira}:D{ultima})", NEGRITO, "0.0000")
    escreve(ws, f"E{tot}", f"=SUM(E{primeira}:E{ultima})", NEGRITO, "0.0000")
    escreve(ws, f"F{tot}", f"=SUM(F{primeira}:F{ultima})", NEGRITO, "0.0%")
    escreve(ws, f"A{tot + 2}", "W (horas) = N / D")
    escreve(ws, f"B{tot + 2}", f"=D{tot}/E{tot}", formato="0.0000")
    escreve(ws, f"A{tot + 3}", "W (minutos)", NEGRITO)
    escreve(ws, f"B{tot + 3}", f"=B{tot + 2}*60", NEGRITO, "0.00")
    escreve(ws, f"A{tot + 5}", "Soma dos tempos = tempo global?")
    escreve(ws, f"B{tot + 5}", f'=IF(ABS(B{tot}-B7)<0.001,"OK","ERRO")')
    larguras(ws, [34, 22, 15, 14, 12, 12])
    nome = ws.title
    return {"perdas": f"'{nome}'!$B$8", "N": f"'{nome}'!$D${tot}",
            "D": f"'{nome}'!$E${tot}", "U": f"'{nome}'!$F${tot}",
            "W": f"'{nome}'!$B${tot + 3}"}


def aba_sementes(wb, execucoes):
    # execucoes: {(modelo, fila): [indices por semente]}. Devolve as médias.
    ws = wb.create_sheet("Sementes")
    escreve(ws, "A1", f"Mesmos modelos com {len(SEEDS)} sementes (100.000 aleatórios cada)", TITULO)
    escreve(ws, "A2", "Azul = resultado de cada simulação. A média e o desvio padrão são fórmulas. "
                      "Mostra que os ganhos da melhoria não dependem da semente escolhida.")
    cabecalho(ws, 4, ["Modelo", "Fila", "Semente", "Perdas", "N", "D (cl/h)", "U", "W (min)"])
    medias = {}
    linha = 5
    for (modelo, fid), lista in execucoes.items():
        primeira = linha
        for seed, ind in zip(SEEDS, lista):
            escreve(ws, f"A{linha}", modelo)
            escreve(ws, f"B{linha}", fid)
            escreve(ws, f"C{linha}", seed)
            escreve(ws, f"D{linha}", ind["perdas"], ENTRADA, "#,##0")
            escreve(ws, f"E{linha}", ind["N"], ENTRADA, "0.0000")
            escreve(ws, f"F{linha}", ind["D"], ENTRADA, "0.0000")
            escreve(ws, f"G{linha}", ind["U"], ENTRADA, "0.0%")
            escreve(ws, f"H{linha}", ind["W_min"], ENTRADA, "0.00")
            linha += 1
        ultima = linha - 1
        escreve(ws, f"A{linha}", modelo, NEGRITO)
        escreve(ws, f"B{linha}", fid, NEGRITO)
        escreve(ws, f"C{linha}", "Média", NEGRITO)
        escreve(ws, f"A{linha + 1}", modelo)
        escreve(ws, f"B{linha + 1}", fid)
        escreve(ws, f"C{linha + 1}", "Desvio padrão")
        for col, fmt in zip("DEFGH", ["#,##0.0", "0.0000", "0.0000", "0.0%", "0.00"]):
            escreve(ws, f"{col}{linha}", f"=AVERAGE({col}{primeira}:{col}{ultima})", NEGRITO, fmt)
            escreve(ws, f"{col}{linha + 1}", f"=STDEV({col}{primeira}:{col}{ultima})", formato=fmt)
        medias[(modelo, fid)] = {k: f"'Sementes'!${col}${linha}"
                                 for k, col in zip(["perdas", "N", "D", "U", "W"], "DEFGH")}
        linha += 3
    larguras(ws, [12, 8, 14, 11, 11, 11, 10, 11])
    return medias


def aba_comparacao(wb, refs, medias, filas_ids):
    ws = wb["Comparação"]
    escreve(ws, "A1", "Comparação: modelo atual x modelo melhorado", TITULO)
    escreve(ws, "A2", "Verde = ligação com as outras abas. Variação = melhorado / atual - 1. "
                      "Tempo em minutos, vazão em clientes por hora.")
    cabecalho(ws, 4, ["Fila", "Índice", "Atual (sem. 1)", "Melhorado (sem. 1)", "Variação",
                      f"Atual (média {len(SEEDS)} sem.)", f"Melhorado (média {len(SEEDS)} sem.)",
                      "Variação"])
    indices = [("perdas", "Perdas", "#,##0.0"), ("N", "População (N)", "0.00"),
               ("D", "Vazão (D)", "0.00"), ("U", "Utilização (U)", "0.0%"),
               ("W", "Tempo de resposta (W)", "0.0")]
    linha = 5
    for fid in filas_ids:
        for chave, rotulo, fmt in indices:
            escreve(ws, f"A{linha}", f"{fid} ({NOMES.get(fid, fid)})")
            escreve(ws, f"B{linha}", rotulo)
            fmt_sem1 = "#,##0" if chave == "perdas" else fmt  # uma simulação: perdas inteiras
            escreve(ws, f"C{linha}", f"={refs[('Atual', fid)][chave]}", LIGACAO, fmt_sem1)
            escreve(ws, f"D{linha}", f"={refs[('Melhorado', fid)][chave]}", LIGACAO, fmt_sem1)
            escreve(ws, f"E{linha}", f'=IF(C{linha}=0,"-",D{linha}/C{linha}-1)',
                    formato="+0.0%;-0.0%;0.0%").alignment = DIREITA
            escreve(ws, f"F{linha}", f"={medias[('Atual', fid)][chave]}", LIGACAO, fmt)
            escreve(ws, f"G{linha}", f"={medias[('Melhorado', fid)][chave]}", LIGACAO, fmt)
            escreve(ws, f"H{linha}", f'=IF(F{linha}=0,"-",G{linha}/F{linha}-1)',
                    formato="+0.0%;-0.0%;0.0%").alignment = DIREITA
            linha += 1
        linha += 1
    escreve(ws, f"A{linha}", "A vazão (D) do N2 e do N3 cai porque, com menos chamados escalando, "
                             "chega menos trabalho a eles; não é perda de capacidade.")
    larguras(ws, [24, 24, 16, 18, 11, 20, 22, 11])


def main():
    parser = argparse.ArgumentParser(description="Gera a planilha de índices do T2")
    parser.add_argument("--atual", default="modelo_atual.yml")
    parser.add_argument("--melhorado", default="modelo_melhorado.yml")
    parser.add_argument("--saida", default="indices_desempenho.xlsx")
    args = parser.parse_args()
    caminhos = {"Atual": args.atual, "Melhorado": args.melhorado}

    modelos, simulacoes = {}, {}
    for nome, caminho in caminhos.items():
        filas, aleatorios = carregar_modelo(caminho)
        modelos[nome] = filas
        simulacoes[nome] = simular_rede(filas, aleatorios, seed=1)

    wb = Workbook()
    wb.active.title = "Comparação"
    parametros = aba_parametros(wb, modelos)
    refs = {}
    for nome, filas in modelos.items():
        sim = simulacoes[nome]
        for f in filas:
            refs[(nome, f.id)] = aba_fila(wb, nome, f, sim["tempo_global"],
                                          sim["filas"][f.id], parametros[(nome, f.id)])

    execucoes = {}
    for nome, caminho in caminhos.items():
        por_seed = [analisar(caminho, s)[0] for s in SEEDS]
        for f in modelos[nome]:
            execucoes[(nome, f.id)] = [r[f.id] for r in por_seed]
    medias = aba_sementes(wb, execucoes)

    aba_comparacao(wb, refs, medias, [f.id for f in modelos["Atual"]])
    # ordem das abas: Comparação, Parâmetros, filas do atual, filas do melhorado, Sementes
    for ws in wb.worksheets:  # impressão: paisagem, cabendo na largura da página
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.save(args.saida)
    print(f"Planilha gerada: {args.saida}")


if __name__ == "__main__":
    main()
